"""Build a 5-slide Thai summary deck of Kise et al. (BMJ 2016) as .pptx (importable into Canva)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
FONT = "Sarabun"

NAVY = RGBColor(0x1F, 0x3B, 0x57)
TEAL = RGBColor(0x14, 0xA3, 0x8B)
CORAL = RGBColor(0xE8, 0x6A, 0x50)
LIGHT = RGBColor(0xEE, 0xF6, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x65, 0x75)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, lines, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: str or list of str / (str, dict) tuples; dict may override size, color, bold."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        opts = {}
        if isinstance(ln, tuple):
            ln, opts = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opts.get("after", 4))
        r = p.add_run(); r.text = ln
        r.font.name = FONT
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", bold)
        r.font.color.rgb = opts.get("color", color)
    return tb


def header(slide, num, title):
    box(slide, 0, 0, 13.333, 1.05, NAVY, MSO_SHAPE.RECTANGLE)
    box(slide, 0.45, 0.22, 0.62, 0.62, TEAL, MSO_SHAPE.OVAL)
    text(slide, 0.45, 0.22, 0.62, 0.62, str(num), 22, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, 1.25, 0.12, 11.5, 0.85, title, 28, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.45, 7.05, 12.4, 0.35,
         "Kise NJ, et al. BMJ 2016;354:i3740 (republished in Br J Sports Med)  |  OMEX trial NCT01002794",
         11, GREY)


def card(slide, x, y, w, h, title, body, accent=TEAL, size=15):
    box(slide, x, y, w, h, LIGHT)
    box(slide, x, y, 0.12, h, accent, MSO_SHAPE.RECTANGLE)
    text(slide, x + 0.25, y + 0.12, w - 0.4, 0.5, title, 18, accent, True)
    text(slide, x + 0.25, y + 0.62, w - 0.4, h - 0.7, body, size, NAVY)


# ---------- Slide 1: Cover ----------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY, MSO_SHAPE.RECTANGLE)
box(s, 0, 6.1, 13.333, 1.4, TEAL, MSO_SHAPE.RECTANGLE)
s.shapes.add_picture(os.path.join(IMG, "knee.png"), Inches(10.1), Inches(1.1), height=Inches(3.6))
text(s, 0.8, 0.7, 9, 0.5, "สรุปงานวิจัย  |  RANDOMISED CONTROLLED TRIAL", 16, TEAL, True)
text(s, 0.8, 1.25, 9.2, 2.4,
     "การออกกำลังกายเทียบกับการผ่าตัดส่องกล้องตัดหมอนรองข้อเข่าบางส่วน "
     "ในผู้ป่วยวัยกลางคนที่หมอนรองข้อเข่าฉีกขาดแบบเสื่อม: ติดตามผล 2 ปี", 32, WHITE, True)
text(s, 0.8, 3.75, 9.2, 1.2,
     "Exercise therapy versus arthroscopic partial meniscectomy for degenerative meniscal tear "
     "in middle aged patients: randomised controlled trial with two year follow-up", 16,
     RGBColor(0xC9, 0xD8, 0xE4))
text(s, 0.8, 4.95, 9.2, 0.9,
     ["Kise NJ, Risberg MA, Stensrud S, Ranstam J, Engebretsen L, Roos EM",
      "BMJ 2016;354:i3740  •  Open Access"], 15, RGBColor(0xC9, 0xD8, 0xE4))
text(s, 0.8, 6.35, 11.8, 0.9, "ผู้นำเสนอ: ..............................................   รหัสนักศึกษา: ..............................",
     20, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)

# ---------- Slide 2: Introduction + Methods ----------
s = prs.slides.add_slide(BLANK)
header(s, 2, "บทนำ และวิธีการวิจัย")
card(s, 0.45, 1.3, 6.1, 2.65, "บทนำ (สรุป)", [
    "• หมอนรองข้อเข่าฉีกขาดแบบเสื่อม (degenerative meniscal tear) พบบ่อยในวัยกลางคน",
    "• การผ่าตัดส่องกล้องตัดหมอนรองบางส่วน (APM) เป็นการผ่าตัดกระดูกที่ทำบ่อยที่สุดอย่างหนึ่ง แต่หลักฐานว่าดีกว่าการไม่ผ่าตัดยังไม่ชัด",
])
card(s, 0.45, 4.1, 6.1, 2.8, "วัตถุประสงค์", [
    "• เปรียบเทียบผลของการออกกำลังกายภายใต้การดูแล 12 สัปดาห์ กับการผ่าตัด APM ต่อผลลัพธ์ที่ผู้ป่วยรายงานเอง (KOOS) ที่ 2 ปี",
    "• และเปรียบเทียบความแข็งแรงกล้ามเนื้อต้นขาที่ 3 และ 12 เดือน",
], CORAL)
card(s, 6.8, 1.3, 6.1, 5.6, "วิธีการวิจัย", [
    "• รูปแบบ: RCT แบบ 2 กลุ่ม, superiority trial, ประเทศนอร์เวย์",
    "• ผู้เข้าร่วม: 140 คน อายุ 35–60 ปี (เฉลี่ย ~49 ปี) หมอนรองฉีกขาดแบบเสื่อมยืนยันด้วย MRI ส่วนใหญ่ไม่มีข้อเข่าเสื่อมชัดเจนในภาพเอกซเรย์",
    "• สุ่มแบ่งกลุ่ม 1:1",
    ("   กลุ่มที่ 1  ออกกำลังกาย 12 สัปดาห์ (n = 70)", {"bold": True, "color": TEAL}),
    ("   กลุ่มที่ 2  ผ่าตัด APM (n = 70)", {"bold": True, "color": CORAL}),
    "• ผลลัพธ์หลัก: การเปลี่ยนแปลงคะแนน KOOS4 (ปวด, อาการ, กีฬา/นันทนาการ, คุณภาพชีวิต) จากเริ่มต้นถึง 2 ปี",
    "• ผลลัพธ์รอง: คะแนน KOOS รายด้าน, ความแข็งแรงกล้ามเนื้อต้นขา และการทดสอบสมรรถภาพ ที่ 3 และ 12 เดือน",
], NAVY, 15)

# ---------- Slide 3: Equipment / tools ----------
s = prs.slides.add_slide(BLANK)
header(s, 3, "อุปกรณ์และเครื่องมือที่ใช้ในงานวิจัย")
tools = [
    ("MRI และเอกซเรย์", "ยืนยันการฉีกขาดของหมอนรองข้อเข่า และประเมินระดับข้อเข่าเสื่อมก่อนเข้าร่วม", NAVY),
    ("แบบสอบถาม KOOS", "Knee injury and Osteoarthritis Outcome Score วัดอาการปวด อาการ การใช้งาน และคุณภาพชีวิต", TEAL),
    ("Isokinetic dynamometer", "เครื่องวัดความแข็งแรงกล้ามเนื้อเหยียดเข่า (quadriceps) และงอเข่า (hamstrings)", CORAL),
    ("การทดสอบสมรรถภาพ", "กระโดดขาเดียวไกล (one-leg hop), กระโดดขาเดียว 6 เมตรจับเวลา (6-m timed hop), ย่อเข่าใน 30 วินาที (knee bend)", NAVY),
    ("อุปกรณ์การฝึก", "จักรยานอยู่กับที่, เครื่อง leg press / knee extension / hamstring curl, น้ำหนัก, กล่องสเต็ป, แผ่นทรงตัว", TEAL),
    ("กลุ่มผ่าตัด", "การผ่าตัดส่องกล้องตัดหมอนรองข้อเข่าบางส่วน (arthroscopic partial meniscectomy) + คำแนะนำท่าบริหารที่บ้าน", CORAL),
]
for i, (t, b, c) in enumerate(tools):
    x = 0.45 + (i % 3) * 4.2; y = 1.35 + (i // 3) * 2.85
    card(s, x, y, 3.95, 2.6, t, b, c, 15)

# ---------- Slide 4: Exercise program ----------
s = prs.slides.add_slide(BLANK)
header(s, 4, "โปรแกรมการออกกำลังกาย 12 สัปดาห์ (กลุ่มออกกำลังกาย)")
# timeline strip
box(s, 0.45, 1.25, 12.45, 0.95, LIGHT)
text(s, 0.6, 1.3, 12.2, 0.85, [
    ("12 สัปดาห์  •  2–3 ครั้ง/สัปดาห์ (รวม 24–36 ครั้ง)  •  มีนักกายภาพบำบัดดูแล  •  เพิ่มความยากทีละระดับ (progressive)",
     {"bold": True, "size": 17}),
], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

sections = [
    ("1  อบอุ่นร่างกาย", "~10 นาที", NAVY, [("bike", "ปั่นจักรยานอยู่กับที่")]),
    ("2  Neuromuscular", "ทรงตัว + ควบคุมแนวเข่า", TEAL,
     [("squat", "Squat"), ("lunge", "Lunge"), ("stepup", "Step-up"),
      ("singlesquat", "Single-leg squat"), ("balance", "ยืนขาเดียวบนแผ่นทรงตัว")]),
    ("3  Strength", "เพิ่มน้ำหนักทีละขั้น", CORAL,
     [("legpress", "Leg press"), ("kneeext", "Knee extension"), ("hamcurl", "Hamstring curl")]),
]
widths = [2.0, 6.45, 3.85]
x = 0.45
for (title, sub, col, items), w in zip(sections, widths):
    box(s, x, 2.4, w, 4.5, LIGHT)
    box(s, x, 2.4, w, 0.75, col)
    text(s, x, 2.4, w, 0.45, title, 17, WHITE, True, PP_ALIGN.CENTER)
    text(s, x, 2.77, w, 0.35, sub, 12, WHITE, False, PP_ALIGN.CENTER)
    n = len(items)
    cols = min(n, 3); rows = (n + cols - 1) // cols
    cw = w / cols
    for j, (img, label) in enumerate(items):
        cx = x + (j % cols) * cw; cy = 3.25 + (j // cols) * (3.55 / rows)
        ih = min(1.35 if rows > 1 else 2.2, cw - 0.05)
        s.shapes.add_picture(os.path.join(IMG, img + ".png"),
                             Inches(cx + (cw - ih) / 2), Inches(cy), height=Inches(ih))
        text(s, cx, cy + ih, cw, 0.4, label, 13, NAVY, True, PP_ALIGN.CENTER)
    x += w
for sx in (0.45 + widths[0], 0.45 + widths[0] + widths[1]):
    box(s, sx - 0.03, 2.4, 0.06, 4.5, WHITE, MSO_SHAPE.RECTANGLE)
text(s, 0.45 + widths[0] + widths[1] + 0.15, 5.0, widths[2] - 0.3, 1.8,
     [("2–3 เซต/ท่า", {"bold": True, "size": 16, "color": CORAL}),
      ("ทำซ้ำประมาณ 8–15 ครั้ง/เซต", {"size": 14}),
      ("เพิ่มน้ำหนักเมื่อทำได้ครบอย่างถูกต้อง", {"size": 14})], 14, NAVY, False, PP_ALIGN.CENTER)

# ---------- Slide 5: Results + Discussion ----------
s = prs.slides.add_slide(BLANK)
header(s, 5, "ผลการวิจัย และอภิปรายผล")
# KPI tiles
kpis = [
    ("0.9", "คะแนน", "ความต่าง KOOS4 ที่ 2 ปี\n95% CI −4.3 ถึง 6.1\nP = 0.72\nไม่แตกต่างกันทางคลินิก", NAVY),
    ("↑", "แรง", "ความแข็งแรงกล้ามเนื้อต้นขา\nกลุ่มออกกำลังกายดีกว่า\nที่ 3 และ 12 เดือน", TEAL),
    ("19%", "", "ของกลุ่มออกกำลังกาย\n(13 จาก 70 คน)\nเปลี่ยนไปผ่าตัดภายใน 2 ปี", CORAL),
]
for i, (big, unit, cap, col) in enumerate(kpis):
    x = 0.45 + i * 2.2
    box(s, x, 1.3, 2.05, 3.05, col)
    text(s, x, 1.4, 2.05, 1.0, big, 40, WHITE, True, PP_ALIGN.CENTER)
    text(s, x + 0.08, 2.45, 1.9, 1.85, cap.split("\n"), 12, WHITE, False, PP_ALIGN.CENTER)
card(s, 0.45, 4.5, 6.45, 2.4, "ผลการวิจัย (สรุป)", [
    "• ทั้ง 2 กลุ่มดีขึ้นใกล้เคียงกันที่ 2 ปี",
    "• กลุ่มออกกำลังกายทำ 6-m timed hop ได้ดีกว่าที่ 3 เดือน (P = 0.02) และ 12 เดือน (P = 0.04)",
    "• ไม่พบเหตุการณ์ไม่พึงประสงค์ร้ายแรงในกลุ่มออกกำลังกาย",
], TEAL, 14)
card(s, 7.1, 1.3, 5.8, 5.6, "อภิปรายผล", [
    "• การออกกำลังกายแบบมีผู้ดูแลให้ผลต่ออาการปวดและการใช้งานไม่ด้อยกว่าการผ่าตัด",
    "• ได้ประโยชน์เพิ่มเรื่องความแข็งแรงกล้ามเนื้อต้นขา ซึ่งเป็นปัจจัยป้องกันข้อเข่าเสื่อม",
    "• หลีกเลี่ยงความเสี่ยงและค่าใช้จ่ายของการผ่าตัด",
    "• ข้อจำกัด: ปกปิดกลุ่มการรักษาจากผู้ป่วยไม่ได้, มีผู้เปลี่ยนกลุ่มไปผ่าตัด",
    ("สรุป: ควรพิจารณาการออกกำลังกายภายใต้การดูแลเป็นทางเลือกแรก ก่อนผ่าตัด ในผู้ป่วยวัยกลางคนที่หมอนรองฉีกขาดแบบเสื่อมและไม่มีข้อเข่าเสื่อมชัดเจน",
     {"bold": True, "color": CORAL}),
], NAVY, 15)

out = os.path.join(HERE, "Meniscus_Exercise_vs_Surgery.pptx")
prs.save(out)
print(out)
